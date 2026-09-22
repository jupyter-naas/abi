from __future__ import annotations

import io
import re

from naas_abi.apps.nexus.sheets.model import SheetWorkbook

_PCT_HEADER = re.compile(r"%|percent|margin|ratio|variance\s*%", re.IGNORECASE)
_MONEY_HEADER = re.compile(
    r"eur|usd|cash|amount|revenue|cogs|opex|payroll|budget|actual|inflow|outflow|cost",
    re.IGNORECASE,
)


def _header_format(header: str) -> str | None:
    text = (header or "").strip()
    if not text:
        return None
    if _PCT_HEADER.search(text):
        return "0.0%"
    if _MONEY_HEADER.search(text):
        return "#,##0.0;(#,##0.0);-"
    return None


def workbook_to_xlsx_bytes(workbook: SheetWorkbook) -> bytes:
    """Export the live model to XLSX, preserving ``=`` formula strings."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("XLSX export requires openpyxl") from exc

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    input_font = Font(color="0000FF", name="Arial", size=11)
    formula_font = Font(color="000000", name="Arial", size=11)
    header_font = Font(bold=True, name="Arial", size=11)
    input_fill = PatternFill("solid", fgColor="FFFF99")

    for tab in workbook.sheets or [SheetWorkbook().sheets[0]]:
        ws = wb.create_sheet(title=(tab.name or "Sheet1")[:31])
        for column, pixels in tab.column_widths.items():
            # Excel column widths use character units; rows use points (96 dpi).
            ws.column_dimensions[openpyxl.utils.get_column_letter(column + 1)].width = (
                pixels - 5
            ) / 7
        for row, pixels in tab.row_heights.items():
            ws.row_dimensions[row + 1].height = pixels * 0.75
        headers = tab.rows[0] if tab.rows else []
        col_formats: dict[int, str] = {}
        for c_idx, header in enumerate(headers):
            fmt = _header_format(str(header) if header is not None else "")
            if fmt:
                col_formats[c_idx] = fmt

        for r_idx, row in enumerate(tab.rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if r_idx == 1:
                    cell.font = header_font
                    continue
                is_formula = isinstance(value, str) and value.strip().startswith("=")
                is_input = (
                    not is_formula
                    and isinstance(value, (int, float))
                    and tab.name.lower() in {"assumptions", "drivers", "inputs"}
                    and c_idx > 1
                )
                if is_formula:
                    cell.font = formula_font
                elif is_input and isinstance(value, (int, float)):
                    cell.font = input_font
                    cell.fill = input_fill
                else:
                    cell.font = formula_font
                fmt = col_formats.get(c_idx - 1)
                if fmt and isinstance(value, (int, float)) or fmt and is_formula:
                    cell.number_format = fmt

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
