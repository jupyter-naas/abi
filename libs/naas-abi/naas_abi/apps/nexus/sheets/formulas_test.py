from pathlib import Path

from naas_abi.apps.nexus.sheets.formulas import (
    evaluate_workbook_formulas,
    formula_errors,
)
from naas_abi.apps.nexus.sheets.html_io import parse_workbook_html
from naas_abi.apps.nexus.sheets.model import SheetTab, SheetWorkbook
from naas_abi.apps.nexus.sheets.xlsx_export import workbook_to_xlsx_bytes

_TEMPLATES = Path(__file__).resolve().parents[1] / "assets" / "sheets" / "templates"


def test_evaluate_simple_arithmetic() -> None:
    wb = SheetWorkbook(
        title="Test",
        sheets=[
            SheetTab(
                name="Sheet1",
                rows=[
                    ["10", "20", "=A1+B1"],
                ],
            ),
        ],
    )
    out = evaluate_workbook_formulas(wb)
    assert out.sheets[0].rows[0][2] == 30.0
    # Source formulas preserved on the original.
    assert wb.sheets[0].rows[0][2] == "=A1+B1"


def test_evaluate_cell_reference() -> None:
    wb = SheetWorkbook(
        title="Test",
        sheets=[
            SheetTab(
                name="Sheet1",
                rows=[
                    ["5", "3", "=A1*B1"],
                ],
            ),
        ],
    )
    out = evaluate_workbook_formulas(wb)
    assert out.sheets[0].rows[0][2] == 15.0


def test_evaluate_nested_formula_refs() -> None:
    wb = SheetWorkbook(
        title="Test",
        sheets=[
            SheetTab(
                name="Sheet1",
                rows=[
                    ["Revenue", 100, 40],
                    ["Total", "=B1+C1", None],
                    ["Margin", "=B2-10", None],
                ],
            ),
        ],
    )
    out = evaluate_workbook_formulas(wb)
    assert out.sheets[0].rows[1][1] == 140.0
    assert out.sheets[0].rows[2][1] == 130.0


def test_sum_if_abs_round_and_sheet_refs() -> None:
    wb = SheetWorkbook(
        title="Funcs",
        sheets=[
            SheetTab(
                name="Assumptions",
                rows=[
                    ["x", 10],
                    ["y", 2.4],
                ],
            ),
            SheetTab(
                name="P&L",
                rows=[
                    ["a", 1, 2, 3],
                    ["sum", "=SUM(B1:D1)", None, None],
                    ["from_assump", "=Assumptions!B1", None, None],
                    ["rounded", "=ROUND(Assumptions!B2,0)", None, None],
                    ["flag", '=IF(ABS(B2-6)<0.001,"OK","BREAK")', None, None],
                ],
            ),
        ],
    )
    out = evaluate_workbook_formulas(wb)
    assert out.sheets[1].rows[1][1] == 6.0
    assert out.sheets[1].rows[2][1] == 10.0
    assert out.sheets[1].rows[3][1] == 2.0
    assert out.sheets[1].rows[4][1] == "OK"
    assert formula_errors(wb) == []


def _assert_finance_template(stem: str, expected_tabs: list[str]) -> None:
    html = (_TEMPLATES / f"{stem}.html").read_text(encoding="utf-8")
    wb = parse_workbook_html(html)
    assert [t.name for t in wb.sheets] == expected_tabs
    assert any(t.name == "Assumptions" for t in wb.sheets)
    assert any(t.name == "Checks" for t in wb.sheets)

    # Tie-out / total cells must be formulas, not hardcoded.
    for tab in wb.sheets:
        if tab.name == "Assumptions":
            continue
        for row in tab.rows[1:]:
            label = str(row[0] or "").lower()
            if any(
                key in label
                for key in (
                    "total",
                    "gross",
                    "ebitda",
                    "operating",
                    "variance",
                    "net change",
                    "closing",
                    "runway",
                    "h1 cost",
                    "check",
                    "diff",
                    "status",
                    "all checks",
                )
            ):
                for cell in row[1:]:
                    if cell in (None, ""):
                        continue
                    if isinstance(cell, str) and cell.startswith("="):
                        continue
                    # Labels / month names on the left are fine; numeric literals
                    # on calc tabs for these rows are not.
                    if isinstance(cell, (int, float)) and tab.name != "Assumptions":
                        # Expected count column on Checks is an intentional literal.
                        if tab.name == "Checks" and "OK count" in str(row[0]):
                            continue
                        raise AssertionError(
                            f"{stem} {tab.name} row {row[0]!r} has literal {cell!r}"
                        )

    errs = formula_errors(wb)
    assert errs == [], errs
    evaluated = evaluate_workbook_formulas(wb)
    checks = next(t for t in evaluated.sheets if t.name == "Checks")
    for row in checks.rows[1:]:
        status = row[4] if len(row) > 4 else None
        assert status == "OK", f"{stem} check {row[0]!r} status={status!r} row={row}"

    # Export keeps formula strings (openpyxl stores them as formulas).
    xlsx = workbook_to_xlsx_bytes(wb)
    assert xlsx[:2] == b"PK"
    import openpyxl

    book = openpyxl.load_workbook(filename := __import__("io").BytesIO(xlsx))
    del filename
    sample = None
    for name in book.sheetnames:
        ws = book[name]
        for row in ws.iter_rows(min_row=2, max_row=30, max_col=8):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    sample = cell.value
                    break
            if sample:
                break
        if sample:
            break
    assert sample and sample.startswith("=")


def test_monthly_pnl_tieouts() -> None:
    _assert_finance_template("monthly-pnl-v1", ["Assumptions", "P&L", "Checks"])


def test_budget_vs_actuals_tieouts() -> None:
    _assert_finance_template(
        "budget-vs-actuals-v1",
        ["Assumptions", "Summary", "Headcount", "Checks"],
    )


def test_cash_runway_tieouts() -> None:
    _assert_finance_template(
        "cash-runway-v1",
        ["Assumptions", "Runway", "Hiring plan", "Checks"],
    )


def test_xlsx_export_preserves_pnl_formulas() -> None:
    html = (_TEMPLATES / "monthly-pnl-v1.html").read_text(encoding="utf-8")
    wb = parse_workbook_html(html)
    assert wb.sheets[1].rows[7][1] == "=B4-B7"
    xlsx = workbook_to_xlsx_bytes(wb)
    import io

    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert book["P&L"]["B8"].value == "=B4-B7"


def test_arithmetic_precedence_and_unsupported_python_expressions() -> None:
    workbook = SheetWorkbook(
        sheets=[
            SheetTab(
                rows=[
                    [
                        "=-(2+3)*4/2",
                        "=2+3*4",
                        "=2**1000000",
                        "=4//2",
                        "=__import__('os').getcwd()",
                        "=(1).__class__",
                        "=1/0",
                    ]
                ]
            )
        ]
    )
    result = evaluate_workbook_formulas(workbook)
    assert result.sheets[0].rows[0] == [-10, 14, "#ERR", "#ERR", "#ERR", "#ERR", "#ERR"]
    assert workbook.sheets[0].rows[0][0] == "=-(2+3)*4/2"


def test_round_matches_spreadsheet_ties_and_negative_digits():
    wb = SheetWorkbook(
        title="round",
        sheets=[
            SheetTab(
                name="S",
                rows=[
                    [
                        "=ROUND(2.5,0)",
                        "=ROUND(-2.5,0)",
                        "=ROUND(1.005,2)",
                        "=ROUND(-1.005,2)",
                        "=ROUND(125,-1)",
                        "=ROUND(-125,-1)",
                    ]
                ],
            )
        ],
    )
    assert evaluate_workbook_formulas(wb).sheets[0].rows[0] == [
        3,
        -3,
        1.01,
        -1.01,
        130,
        -130,
    ]


def test_reference_precision_and_small_values_are_preserved():
    wb = SheetWorkbook(
        title="precision",
        sheets=[
            SheetTab(
                name="S",
                rows=[
                    [
                        123456789012345,
                        "=A1+1",
                        0.000000000000001,
                        "=C1*1000",
                    ]
                ],
            )
        ],
    )
    row = evaluate_workbook_formulas(wb).sheets[0].rows[0]
    assert row[1] == 123456789012346
    assert row[3] == 1e-12


def test_repeated_dependencies_are_memoized(monkeypatch):
    from naas_abi.apps.nexus.sheets import formulas

    calls = 0
    original = formulas._eval_expr

    def counted(*args, **kwargs):
        nonlocal calls
        calls += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(formulas, "_eval_expr", counted)
    rows = [[1]] + [[f"=A{i}+A{i}"] for i in range(1, 35)]
    wb = SheetWorkbook(title="fanout", sheets=[SheetTab(name="S", rows=rows)])
    assert evaluate_workbook_formulas(wb).sheets[0].rows[-1][0] == 2**34
    assert calls == 34


def test_workbook_budget_is_shared_and_resets_between_requests(monkeypatch):
    from naas_abi.apps.nexus.sheets import formulas

    original = formulas._Evaluation
    monkeypatch.setattr(formulas, "_Evaluation", lambda: original(remaining=40))
    wb = SheetWorkbook(
        title="budget", sheets=[SheetTab(name="S", rows=[["=1+2"] for _ in range(30)])]
    )
    calculated, errors = formulas.calculate_workbook(wb)
    assert calculated.sheets[0].rows[0][0] == 3
    assert errors and "limit exceeded" in errors[-1]["error"]
    assert formulas.calculate_workbook(wb)[0].sheets[0].rows[0][0] == 3
